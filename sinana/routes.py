import math
import os
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from uuid import uuid4

from flask import Blueprint, abort, current_app, flash, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from flask_login import current_user, login_required, login_user, logout_user
from openpyxl import load_workbook
from sqlalchemy import func
from werkzeug.utils import secure_filename

from . import db
from .models import (
    Counterparty,
    Customer,
    LocationPing,
    Order,
    OrderItem,
    PhotoReport,
    PriceType,
    Product,
    ProductPrice,
    ProductionOrder,
    ReceivableTransaction,
    Recipe,
    RecipeItem,
    SalesReturn,
    SalesReturnItem,
    Setting,
    StockBalance,
    SyncLog,
    User,
    Visit,
    Warehouse,
)

auth = Blueprint("auth", __name__)
main = Blueprint("main", __name__)
api = Blueprint("api", __name__)
PHOTO_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "heic"}
ORDER_STATUSES = {"submitted", "confirmed", "assembled", "delivered", "paid", "cancelled"}
UNIT_ALIASES = {
    "г": "г", "гр": "г", "грамм": "г", "граммы": "г", "g": "г",
    "кг": "кг", "килограмм": "кг", "килограммы": "кг", "kg": "кг",
    "шт": "шт", "шт.": "шт", "штука": "шт", "штуки": "шт", "pcs": "шт",
}


def manager_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if not current_user.is_manager:
            abort(403)
        return view(*args, **kwargs)

    return wrapped


def one_c_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        configured = current_app.config.get("ONE_C_API_TOKEN") or ""
        supplied = request.headers.get("X-API-Key") or request.headers.get("Authorization", "").removeprefix("Bearer ")
        if not configured:
            return jsonify(error="Интеграция с 1С не настроена"), 503
        if not supplied or supplied != configured:
            return jsonify(error="Недействительный ключ"), 401
        return view(*args, **kwargs)

    return wrapped


def setting(key, default=""):
    row = db.session.get(Setting, key)
    return row.value if row else default


def bool_setting(key, default=False):
    return setting(key, "1" if default else "0") == "1"


def decimal_value(value, default=Decimal("0")):
    try:
        return Decimal(str(value or "0").replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return default


def barcode_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_unit(value):
    return UNIT_ALIASES.get(str(value or "шт").strip().lower(), "шт")


def convert_quantity(quantity, from_unit, to_unit):
    from_unit, to_unit = normalize_unit(from_unit), normalize_unit(to_unit)
    if from_unit == to_unit:
        return quantity
    if from_unit == "г" and to_unit == "кг":
        return quantity / Decimal("1000")
    if from_unit == "кг" and to_unit == "г":
        return quantity * Decimal("1000")
    raise ValueError("Нельзя пересчитать массу в штуки или наоборот")


def parse_coordinate(value):
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def can_manage_visit(visit):
    return current_user.is_manager or visit.representative_id == current_user.id


@main.route("/sw.js")
def service_worker():
    response = send_from_directory(current_app.static_folder, "sw.js")
    response.headers["Service-Worker-Allowed"] = "/"
    response.headers["Cache-Control"] = "no-cache"
    return response


@auth.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard" if current_user.is_manager else "main.mobile"))
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        user = User.query.filter_by(username=username).first()
        if not user or not user.is_active_user or not user.check_password(request.form.get("password") or ""):
            flash("Неверный логин или пароль", "danger")
        else:
            login_user(user)
            session["user_session_version"] = user.session_version
            return redirect(url_for("main.dashboard" if user.is_manager else "main.mobile"))
    return render_template("login.html")


@auth.route("/logout", methods=["POST"])
@login_required
def logout():
    logout_user()
    session.clear()
    return redirect(url_for("auth.login"))


@main.route("/")
@login_required
def dashboard():
    if not current_user.is_manager:
        return redirect(url_for("main.mobile"))
    start = datetime.combine(date.today().replace(day=1), datetime.min.time())
    orders = Order.query.filter(Order.created_at >= start, Order.status != "cancelled").all()
    sales = sum(float(order.total_amount or 0) for order in orders)
    monthly_plan = float(setting("monthly_sales_plan", "1000000") or 1000000)
    plan_percent = round(min(sales / monthly_plan * 100, 999), 1) if monthly_plan else 0
    today = date.today()
    chart_values = []
    for offset in range(4, -1, -1):
        period_end = datetime.combine(today, datetime.max.time())
        period_start = period_end.replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=offset * 7 + 6)
        period_finish = period_start + timedelta(days=7)
        chart_values.append(sum(float(order.total_amount or 0) for order in orders if period_start <= order.created_at < period_finish))
    chart_max = max(chart_values + [monthly_plan / 4 if monthly_plan else 1, 1])
    chart_points = " ".join(f"{40 + index * 155},{190 - value / chart_max * 145:.1f}" for index, value in enumerate(chart_values))
    return render_template(
        "dashboard.html",
        sales=sales,
        orders=orders,
        active_customers=Customer.query.filter_by(is_active=True).count(),
        debt=db.session.query(func.coalesce(func.sum(Counterparty.debt), 0)).scalar(),
        today_visits=Visit.query.filter_by(visit_date=date.today()).all(),
        representatives=User.query.filter_by(role="representative", is_active_user=True).all(),
        unsynced_orders=Order.query.filter(Order.synced_at.is_(None), Order.status != "cancelled").count(),
        monthly_plan=monthly_plan,
        plan_percent=plan_percent,
        chart_values=chart_values,
        chart_points=chart_points,
    )


@main.route("/mobile")
@login_required
def mobile():
    rep = current_user
    if current_user.is_manager and request.args.get("rep_id", type=int):
        rep = User.query.filter_by(id=request.args.get("rep_id", type=int), role="representative").first_or_404()
    visits = Visit.query.filter_by(visit_date=date.today(), representative_id=rep.id).order_by(Visit.sequence).all()
    day_start = datetime.combine(date.today(), datetime.min.time())
    orders = Order.query.filter(Order.representative_id == rep.id, Order.created_at >= day_start).all()
    customers = Customer.query.filter_by(assigned_rep_id=rep.id, is_active=True).order_by(Customer.name).all()
    recent_orders = Order.query.filter_by(representative_id=rep.id).filter(Order.status != "cancelled").order_by(Order.created_at.desc()).limit(30).all()
    return render_template("mobile.html", representative=rep, visits=visits, orders=orders, customers=customers, recent_orders=recent_orders, photo_required=bool_setting("photo_required", True), is_own_route=rep.id == current_user.id)


@main.route("/gps/ping", methods=["POST"])
@login_required
def gps_ping():
    payload = request.get_json(silent=True) or request.form
    latitude = parse_coordinate(payload.get("latitude"))
    longitude = parse_coordinate(payload.get("longitude"))
    accuracy = parse_coordinate(payload.get("accuracy"))
    if latitude is None or longitude is None or not (-90 <= latitude <= 90) or not (-180 <= longitude <= 180):
        return jsonify(error="Некорректные координаты"), 400
    now = datetime.utcnow()
    current_user.last_latitude, current_user.last_longitude, current_user.last_location_at = latitude, longitude, now
    db.session.add(LocationPing(user_id=current_user.id, latitude=latitude, longitude=longitude, accuracy=accuracy, recorded_at=now))
    db.session.commit()
    return jsonify(ok=True, recorded_at=now.isoformat())


@main.route("/visits/<int:visit_id>/start", methods=["POST"])
@login_required
def visit_start(visit_id):
    visit = db.get_or_404(Visit, visit_id)
    if not can_manage_visit(visit):
        abort(403)
    latitude, longitude = parse_coordinate(request.form.get("latitude")), parse_coordinate(request.form.get("longitude"))
    if latitude is None or longitude is None:
        flash("Для первого визита разрешите определение геопозиции", "danger")
        return redirect(url_for("main.mobile"))
    visit.status, visit.started_at = "in_progress", visit.started_at or datetime.utcnow()
    visit.start_latitude, visit.start_longitude = latitude, longitude
    if visit.customer.latitude is None or visit.customer.longitude is None:
        visit.customer.latitude, visit.customer.longitude = latitude, longitude
        visit.customer.gps_captured_at = datetime.utcnow()
        flash("GPS-точка клиента сохранена для следующих маршрутов", "success")
    db.session.commit()
    return redirect(url_for("main.mobile"))


@main.route("/visits/<int:visit_id>/complete", methods=["POST"])
@login_required
def visit_complete(visit_id):
    visit = db.get_or_404(Visit, visit_id)
    if not can_manage_visit(visit):
        abort(403)
    uploaded = request.files.get("photo")
    if bool_setting("photo_required", True) and not uploaded and not visit.photos:
        flash("Для завершения визита требуется фотография выкладки", "danger")
        return redirect(url_for("main.mobile"))
    if uploaded and uploaded.filename:
        ext = uploaded.filename.rsplit(".", 1)[-1].lower() if "." in uploaded.filename else ""
        if ext not in PHOTO_EXTENSIONS:
            flash("Разрешены фотографии JPG, PNG, WEBP или HEIC", "danger")
            return redirect(url_for("main.mobile"))
        stored_name = f"{visit.id}-{uuid4().hex}.{ext}"
        uploaded.save(os.path.join(current_app.config["UPLOAD_FOLDER"], stored_name))
        db.session.add(PhotoReport(visit_id=visit.id, filename=stored_name, original_name=secure_filename(uploaded.filename)))
    visit.status, visit.completed_at = "completed", datetime.utcnow()
    visit.result = request.form.get("result") if request.form.get("result") in {"order", "no_order", "visited"} else "visited"
    db.session.commit()
    flash("Визит завершён", "success")
    return redirect(url_for("main.mobile"))


@main.route("/route/finish", methods=["POST"])
@login_required
def route_finish():
    visits = Visit.query.filter_by(visit_date=date.today(), representative_id=current_user.id).all()
    incomplete = [visit for visit in visits if visit.status != "completed"]
    missing_photos = [visit for visit in visits if not visit.photos] if bool_setting("photo_required", True) else []
    if incomplete:
        flash(f"Сначала завершите все точки: осталось {len(incomplete)}", "danger")
    elif missing_photos:
        flash(f"Не хватает фотографий по {len(missing_photos)} точкам", "danger")
    else:
        flash("Маршрут завершён, данные и фотоматериалы сохранены", "success")
    return redirect(url_for("main.mobile"))


@main.route("/photos/<path:filename>")
@login_required
def photo(filename):
    report = PhotoReport.query.filter_by(filename=filename).first_or_404()
    if not current_user.is_manager and report.visit.representative_id != current_user.id:
        abort(403)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], filename)


@main.route("/orders/new/<int:customer_id>", methods=["GET", "POST"])
@login_required
def order_new(customer_id):
    customer = db.get_or_404(Customer, customer_id)
    if not current_user.is_manager and customer.assigned_rep_id != current_user.id:
        abort(403)
    products = Product.query.filter_by(is_active=True).order_by(Product.name).all()
    prices = {product.id: product_price(product, customer) for product in products}
    if request.method == "POST":
        selected, total = [], Decimal("0")
        for product in products:
            qty = decimal_value(request.form.get(f"qty_{product.id}"))
            if qty > 0:
                price = prices[product.id]
                selected.append((product, qty, price))
                total += qty * price
        if not selected:
            flash("Добавьте хотя бы один товар", "danger")
        else:
            number = f"SIN-{datetime.now():%Y%m%d}-{uuid4().hex[:6].upper()}"
            order = Order(number=number, customer_id=customer.id, representative_id=current_user.id, total_amount=total)
            db.session.add(order)
            db.session.flush()
            for product, qty, price in selected:
                db.session.add(OrderItem(order_id=order.id, product_id=product.id, quantity=qty, price=price))
            visit = Visit.query.filter_by(customer_id=customer.id, representative_id=current_user.id, visit_date=date.today()).first()
            if visit:
                visit.result = "order"
            db.session.commit()
            flash(f"Заказ {number} создан и ожидает передачи в 1С", "success")
            return redirect(url_for("main.mobile"))
    return render_template("order_form.html", customer=customer, products=products, prices=prices)


@main.route("/orders")
@manager_required
def orders():
    return render_template("orders.html", orders=Order.query.order_by(Order.created_at.desc()).limit(300).all())


@main.route("/team")
@manager_required
def team():
    return render_template("team.html", users=User.query.order_by(User.role, User.full_name).all())


@main.route("/team/create", methods=["POST"])
@manager_required
def team_create():
    username = (request.form.get("username") or "").strip()
    full_name, password = (request.form.get("full_name") or "").strip(), request.form.get("password") or ""
    role = request.form.get("role") or "representative"
    if role not in ("manager", "representative"):
        abort(400)
    if not username or not full_name or len(password) < 8:
        flash("Заполните имя, логин и пароль не короче 8 символов", "danger")
    elif User.query.filter_by(username=username).first():
        flash("Такой логин уже существует", "danger")
    else:
        user = User(username=username, full_name=full_name, role=role)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        flash("Сотрудник добавлен", "success")
    return redirect(url_for("main.team"))


@main.route("/team/<int:user_id>/password", methods=["POST"])
@manager_required
def team_password(user_id):
    user, password = db.get_or_404(User, user_id), request.form.get("password") or ""
    if len(password) < 8:
        flash("Пароль должен содержать не менее 8 символов", "danger")
    else:
        user.set_password(password)
        user.session_version += 1
        db.session.commit()
        flash("Пароль изменён, прежние сессии завершены", "success")
    return redirect(url_for("main.team"))


@main.route("/team/<int:user_id>/toggle", methods=["POST"])
@manager_required
def team_toggle(user_id):
    user = db.get_or_404(User, user_id)
    if user.id == current_user.id:
        flash("Нельзя отключить свою учётную запись", "danger")
    else:
        user.is_active_user = not user.is_active_user
        user.session_version += 1
        db.session.commit()
        flash("Статус сотрудника изменён", "success")
    return redirect(url_for("main.team"))


@main.route("/team/<int:user_id>/delete", methods=["POST"])
@manager_required
def team_delete(user_id):
    user = db.get_or_404(User, user_id)
    if user.id == current_user.id:
        flash("Нельзя удалить свою учётную запись", "danger")
    elif Customer.query.filter_by(assigned_rep_id=user.id).first() or Order.query.filter_by(representative_id=user.id).first():
        user.is_active_user = False
        user.session_version += 1
        db.session.commit()
        flash("Сотрудник отключён; история сохранена", "success")
    else:
        db.session.delete(user)
        db.session.commit()
        flash("Сотрудник удалён", "success")
    return redirect(url_for("main.team"))


@main.route("/team/<int:user_id>/role", methods=["POST"])
@manager_required
def team_role(user_id):
    user = db.get_or_404(User, user_id)
    role = request.form.get("role")
    if role not in {"manager", "representative"}:
        abort(400)
    if user.id == current_user.id and role != "manager":
        flash("Нельзя снять роль руководителя у своей учётной записи", "danger")
    else:
        user.role = role
        user.session_version += 1
        db.session.commit()
        flash("Роль изменена; сотрудник войдёт заново", "success")
    return redirect(url_for("main.team"))


@main.route("/customers", methods=["GET", "POST"])
@manager_required
def customers():
    reps = User.query.filter_by(role="representative", is_active_user=True).order_by(User.full_name).all()
    price_types = PriceType.query.filter_by(is_active=True).order_by(PriceType.name).all()
    if request.method == "POST":
        name, address = (request.form.get("name") or "").strip(), (request.form.get("address") or "").strip()
        rep = db.session.get(User, request.form.get("representative_id", type=int))
        latitude, longitude = parse_coordinate(request.form.get("latitude")), parse_coordinate(request.form.get("longitude"))
        coordinates_supplied = bool((request.form.get("latitude") or "").strip() or (request.form.get("longitude") or "").strip())
        coordinates_valid = latitude is not None and longitude is not None and -90 <= latitude <= 90 and -180 <= longitude <= 180
        if not name or not address or not rep or rep.role != "representative":
            flash("Заполните название, адрес и представителя", "danger")
        elif coordinates_supplied and not coordinates_valid:
            flash("Укажите корректные широту и долготу либо оставьте оба поля пустыми", "danger")
        else:
            db.session.add(Customer(
                name=name,
                address=address,
                assigned_rep_id=rep.id,
                counterparty_id=request.form.get("counterparty_id", type=int),
                price_type_id=request.form.get("price_type_id", type=int),
                latitude=latitude if coordinates_valid else None,
                longitude=longitude if coordinates_valid else None,
                gps_captured_at=datetime.utcnow() if coordinates_valid else None,
            ))
            db.session.commit()
            flash("Торговая точка добавлена", "success")
        return redirect(url_for("main.customers"))
    return render_template("customers.html", customers=Customer.query.order_by(Customer.name).all(), representatives=reps, counterparties=Counterparty.query.filter_by(is_active=True).order_by(Counterparty.name).all(), price_types=price_types)


@main.route("/customers/<int:customer_id>/settings", methods=["POST"])
@manager_required
def customer_settings(customer_id):
    customer = db.get_or_404(Customer, customer_id)
    latitude, longitude = parse_coordinate(request.form.get("latitude")), parse_coordinate(request.form.get("longitude"))
    coordinates_supplied = bool((request.form.get("latitude") or "").strip() or (request.form.get("longitude") or "").strip())
    if coordinates_supplied and (latitude is None or longitude is None or not (-90 <= latitude <= 90 and -180 <= longitude <= 180)):
        flash("Укажите корректные GPS-координаты", "danger")
        return redirect(url_for("main.customers"))
    price_type_id = request.form.get("price_type_id", type=int)
    if price_type_id and not db.session.get(PriceType, price_type_id):
        abort(400, "Неизвестный тип цен")
    customer.price_type_id = price_type_id
    if coordinates_supplied:
        customer.latitude, customer.longitude = latitude, longitude
        customer.gps_captured_at = datetime.utcnow()
    db.session.commit()
    flash("Тип цен и GPS торговой точки сохранены", "success")
    return redirect(url_for("main.customers"))


@main.route("/customers/<int:customer_id>/toggle", methods=["POST"])
@manager_required
def customer_toggle(customer_id):
    customer = db.get_or_404(Customer, customer_id)
    customer.is_active = not customer.is_active
    db.session.commit()
    return redirect(url_for("main.customers"))


@main.route("/routes", methods=["GET", "POST"])
@manager_required
def routes():
    reps = User.query.filter_by(role="representative", is_active_user=True).order_by(User.full_name).all()
    if request.method == "POST":
        rep_id, customer_ids = request.form.get("representative_id", type=int), request.form.getlist("customer_ids", type=int)
        try:
            visit_date = datetime.strptime(request.form.get("visit_date") or "", "%Y-%m-%d").date()
        except ValueError:
            visit_date = None
        rep = db.session.get(User, rep_id)
        if not rep or rep.role != "representative" or not visit_date or not customer_ids:
            flash("Выберите сотрудника, дату и торговые точки", "danger")
        else:
            existing = {row.customer_id for row in Visit.query.filter_by(representative_id=rep.id, visit_date=visit_date).all()}
            sequence = len(existing)
            for customer_id in customer_ids:
                customer = db.session.get(Customer, customer_id)
                if customer and customer.id not in existing:
                    sequence += 1
                    db.session.add(Visit(visit_date=visit_date, sequence=sequence, customer_id=customer.id, representative_id=rep.id))
            db.session.commit()
            flash("Маршрут сохранён", "success")
        return redirect(url_for("main.routes", date=(visit_date or date.today()).isoformat()))
    try:
        route_date = datetime.strptime(request.args.get("date") or date.today().isoformat(), "%Y-%m-%d").date()
    except ValueError:
        route_date = date.today()
    visits = Visit.query.filter_by(visit_date=route_date).order_by(Visit.representative_id, Visit.sequence).all()
    return render_template("routes.html", representatives=reps, customers=Customer.query.filter_by(is_active=True).order_by(Customer.name).all(), visits=visits, route_date=route_date)


@main.route("/routes/<int:user_id>/optimize", methods=["POST"])
@manager_required
def route_optimize(user_id):
    try:
        visit_date = datetime.strptime(request.form.get("visit_date") or "", "%Y-%m-%d").date()
    except ValueError:
        abort(400)
    visits = Visit.query.filter_by(visit_date=visit_date, representative_id=user_id).all()
    located = [v for v in visits if v.customer.latitude is not None and v.customer.longitude is not None]
    missing = [v for v in visits if v not in located]
    user = db.get_or_404(User, user_id)
    current = (user.last_latitude, user.last_longitude) if user.last_latitude is not None else None
    ordered = []
    while located:
        chosen = located[0] if current is None else min(located, key=lambda v: (v.customer.latitude - current[0]) ** 2 + (v.customer.longitude - current[1]) ** 2)
        ordered.append(chosen)
        located.remove(chosen)
        current = (chosen.customer.latitude, chosen.customer.longitude)
    for sequence, visit in enumerate(ordered + sorted(missing, key=lambda v: v.sequence), 1):
        visit.sequence = sequence
    db.session.commit()
    flash("Маршрут перестроен по GPS-точкам", "success")
    return redirect(url_for("main.routes", date=visit_date.isoformat()))


@main.route("/tracking")
@manager_required
def tracking():
    return render_template("tracking.html", representatives=User.query.filter_by(role="representative", is_active_user=True).order_by(User.full_name).all())


@main.route("/products", methods=["GET", "POST"])
@manager_required
def products():
    if request.method == "POST":
        if request.form.get("action") == "manual":
            barcode = barcode_value(request.form.get("barcode"))
            name = (request.form.get("name") or "").strip()
            if not barcode or not name:
                flash("Укажите штрихкод и наименование", "danger")
            elif Product.query.filter_by(barcode=barcode).first():
                flash("Товар с таким штрихкодом уже существует", "danger")
            else:
                product_type = request.form.get("product_type") or "finished"
                if product_type not in {"raw", "semi_finished", "finished", "packaging"}:
                    product_type = "finished"
                product = Product(
                    barcode=barcode,
                    name=name,
                    product_type=product_type,
                    unit=normalize_unit(request.form.get("unit")),
                    shelf_life_days=max(request.form.get("shelf_life_days", type=int) or 0, 0) or None,
                    price=max(decimal_value(request.form.get("price")), 0),
                    calculated_cost=max(decimal_value(request.form.get("calculated_cost")), 0),
                    external_group_id=setting("one_c_product_group_id") or None,
                )
                db.session.add(product)
                db.session.commit()
                flash("Номенклатура добавлена", "success")
            return redirect(url_for("main.products"))
        uploaded = request.files.get("file")
        if not uploaded or not uploaded.filename.lower().endswith(".xlsx"):
            flash("Выберите файл Excel в формате .xlsx", "danger")
        else:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook.active
            headers = {str(cell.value or "").strip().lower(): i for i, cell in enumerate(next(sheet.iter_rows()))}
            barcode_col = next((headers[key] for key in headers if key in {"штрихкод", "barcode", "штрих-код"}), None)
            name_col = next((headers[key] for key in headers if key in {"наименование", "название", "name"}), None)
            price_col = next((headers[key] for key in headers if key in {"цена", "price"}), None)
            cost_col = next((headers[key] for key in headers if key in {"себестоимость", "cost"}), None)
            unit_col = next((headers[key] for key in headers if key in {"единица", "ед. изм.", "unit"}), None)
            type_col = next((headers[key] for key in headers if key in {"тип", "type"}), None)
            if barcode_col is None or name_col is None:
                flash("Нужны колонки «Штрихкод» и «Наименование»", "danger")
            else:
                count = 0
                for row in sheet.iter_rows(values_only=True):
                    barcode, name = barcode_value(row[barcode_col]), str(row[name_col] or "").strip()
                    if not barcode or not name:
                        continue
                    product = Product.query.filter_by(barcode=barcode).first() or Product(barcode=barcode)
                    product.name, product.is_active = name, True
                    if price_col is not None:
                        product.price = decimal_value(row[price_col])
                    if cost_col is not None:
                        product.calculated_cost = decimal_value(row[cost_col])
                    if unit_col is not None:
                        product.unit = normalize_unit(row[unit_col])
                    if type_col is not None and str(row[type_col] or "").strip() in {"raw", "semi_finished", "finished", "packaging"}:
                        product.product_type = str(row[type_col]).strip()
                    product.external_group_id = setting("one_c_product_group_id") or product.external_group_id
                    db.session.add(product)
                    count += 1
                db.session.commit()
                flash(f"Загружено или обновлено позиций: {count}", "success")
        return redirect(url_for("main.products"))
    return render_template("products.html", products=Product.query.order_by(Product.name).limit(1000).all(), last_sync=SyncLog.query.order_by(SyncLog.created_at.desc()).first())


@main.route("/settings", methods=["GET", "POST"])
@manager_required
def settings():
    if request.method == "POST":
        values = {
            "photo_required": "1" if request.form.get("photo_required") else "0",
            "monthly_sales_plan": str(max(decimal_value(request.form.get("monthly_sales_plan")), 0)),
            "one_c_product_group_id": (request.form.get("one_c_product_group_id") or "").strip(),
        }
        for key, value in values.items():
            row = db.session.get(Setting, key) or Setting(key=key)
            row.value = value
            db.session.add(row)
        db.session.commit()
        flash("Настройки сохранены", "success")
        return redirect(url_for("main.settings"))
    return render_template("settings.html", photo_required=bool_setting("photo_required", True), monthly_sales_plan=setting("monthly_sales_plan", "1000000"), product_group_id=setting("one_c_product_group_id"), api_configured=bool(current_app.config.get("ONE_C_API_TOKEN")))


def product_price(product, customer=None):
    price_type_id = None
    if customer:
        price_type_id = customer.price_type_id or (customer.counterparty.price_type_id if customer.counterparty else None)
    if price_type_id:
        row = ProductPrice.query.filter_by(product_id=product.id, price_type_id=price_type_id).first()
        if row:
            return row.price
    return product.price


def recalculate_recipe(recipe):
    material_cost = Decimal("0")
    for item in recipe.items:
        quantity_in_stock_unit = convert_quantity(item.quantity, item.quantity_unit or item.component.unit, item.component.unit)
        quantity_with_loss = quantity_in_stock_unit * (Decimal("1") + item.loss_percent / Decimal("100"))
        material_cost += quantity_with_loss * item.component.calculated_cost
    base = material_cost + recipe.labor_cost + recipe.additional_cost
    total = base * (Decimal("1") + recipe.overhead_percent / Decimal("100"))
    unit_cost = total / recipe.output_quantity if recipe.output_quantity else Decimal("0")
    recipe.product.calculated_cost = unit_cost
    return unit_cost


@main.route("/erp")
@manager_required
def erp():
    planned = ProductionOrder.query.filter(ProductionOrder.status.in_(("planned", "in_progress"))).count()
    low_stock = StockBalance.query.filter(StockBalance.quantity <= 0).count()
    return render_template(
        "erp.html",
        products=Product.query.filter_by(is_active=True).count(),
        counterparties=Counterparty.query.filter_by(is_active=True).count(),
        warehouses=Warehouse.query.filter_by(is_active=True).count(),
        production_planned=planned,
        returns_open=SalesReturn.query.filter(SalesReturn.status != "completed").count(),
        low_stock=low_stock,
    )


@main.route("/counterparties", methods=["GET", "POST"])
@manager_required
def counterparties():
    price_types = PriceType.query.filter_by(is_active=True).order_by(PriceType.name).all()
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("Укажите наименование контрагента", "danger")
        else:
            db.session.add(Counterparty(
                name=name,
                tax_id=(request.form.get("tax_id") or "").strip() or None,
                phone=(request.form.get("phone") or "").strip() or None,
                price_type_id=request.form.get("price_type_id", type=int),
                payment_terms_days=max(request.form.get("payment_terms_days", type=int) or 0, 0),
            ))
            db.session.commit()
            flash("Контрагент добавлен", "success")
        return redirect(url_for("main.counterparties"))
    return render_template(
        "counterparties.html",
        counterparties=Counterparty.query.order_by(Counterparty.name).all(),
        price_types=price_types,
        transactions=ReceivableTransaction.query.order_by(ReceivableTransaction.transaction_date.desc(), ReceivableTransaction.id.desc()).limit(100).all(),
        now_date=date.today().isoformat(),
    )


@main.route("/counterparties/<int:counterparty_id>/terms", methods=["POST"])
@manager_required
def counterparty_terms(counterparty_id):
    counterparty = db.get_or_404(Counterparty, counterparty_id)
    price_type_id = request.form.get("price_type_id", type=int)
    if price_type_id and not db.session.get(PriceType, price_type_id):
        abort(400, "Неизвестный тип цен")
    counterparty.price_type_id = price_type_id
    counterparty.payment_terms_days = max(request.form.get("payment_terms_days", type=int) or 0, 0)
    db.session.commit()
    flash("Условия оплаты и тип цен сохранены", "success")
    return redirect(url_for("main.counterparties"))


@main.route("/counterparties/transactions", methods=["POST"])
@manager_required
def counterparty_transaction():
    counterparty = db.get_or_404(Counterparty, request.form.get("counterparty_id", type=int))
    transaction_type = request.form.get("transaction_type")
    amount = max(decimal_value(request.form.get("amount")), 0)
    if transaction_type not in {"payment", "refund"} or amount <= 0:
        flash("Укажите вид операции и положительную сумму", "danger")
        return redirect(url_for("main.counterparties"))
    current_debt = Decimal(counterparty.debt or 0)
    counterparty.debt = current_debt - amount if transaction_type == "payment" else current_debt + amount
    transaction = ReceivableTransaction(
        counterparty_id=counterparty.id,
        transaction_type=transaction_type,
        amount=amount,
        transaction_date=request.form.get("transaction_date", type=date.fromisoformat) or date.today(),
        note=(request.form.get("note") or "").strip() or None,
        balance_after=counterparty.debt,
        created_by_id=current_user.id,
    )
    db.session.add(transaction)
    db.session.commit()
    message = "Оплата клиента зафиксирована" if transaction_type == "payment" else "Возврат денег клиенту зафиксирован"
    flash(message, "success")
    return redirect(url_for("main.counterparties"))


@main.route("/warehouses", methods=["GET", "POST"])
@manager_required
def warehouses():
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name or Warehouse.query.filter_by(name=name).first():
            flash("Укажите уникальное название склада", "danger")
        else:
            db.session.add(Warehouse(name=name, warehouse_type=request.form.get("warehouse_type") or "finished"))
            db.session.commit()
            flash("Склад добавлен", "success")
        return redirect(url_for("main.warehouses"))
    return render_template("warehouses.html", warehouses=Warehouse.query.order_by(Warehouse.name).all(), balances=StockBalance.query.order_by(StockBalance.warehouse_id, StockBalance.product_id).limit(1000).all())


@main.route("/price-types", methods=["GET", "POST"])
@manager_required
def price_types():
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if name and not PriceType.query.filter_by(name=name).first():
            db.session.add(PriceType(name=name))
            db.session.commit()
            flash("Тип цен добавлен", "success")
        return redirect(url_for("main.price_types"))
    return render_template("price_types.html", price_types=PriceType.query.order_by(PriceType.name).all(), products=Product.query.filter_by(is_active=True).order_by(Product.name).all(), prices=ProductPrice.query.all())


@main.route("/price-types/set", methods=["POST"])
@manager_required
def price_set():
    product = db.get_or_404(Product, request.form.get("product_id", type=int))
    price_type = db.get_or_404(PriceType, request.form.get("price_type_id", type=int))
    row = ProductPrice.query.filter_by(product_id=product.id, price_type_id=price_type.id).first() or ProductPrice(product_id=product.id, price_type_id=price_type.id)
    row.price = max(decimal_value(request.form.get("price")), 0)
    db.session.add(row)
    db.session.commit()
    flash("Цена установлена", "success")
    return redirect(url_for("main.price_types"))


@main.route("/recipes", methods=["GET", "POST"])
@manager_required
def recipes():
    products = Product.query.filter_by(is_active=True).order_by(Product.name).all()
    if request.method == "POST":
        finished = db.get_or_404(Product, request.form.get("product_id", type=int))
        finished.product_type = "finished"
        recipe = Recipe.query.filter_by(product_id=finished.id).first() or Recipe(product_id=finished.id)
        recipe.output_quantity = max(decimal_value(request.form.get("output_quantity"), Decimal("1")), Decimal("0.001"))
        recipe.labor_cost = max(decimal_value(request.form.get("labor_cost")), 0)
        recipe.additional_cost = max(decimal_value(request.form.get("additional_cost")), 0)
        recipe.cost_notes = (request.form.get("cost_notes") or "").strip() or None
        recipe.overhead_percent = max(decimal_value(request.form.get("overhead_percent")), 0)
        db.session.add(recipe)
        db.session.commit()
        flash("Комплектация сохранена; добавьте компоненты", "success")
        return redirect(url_for("main.recipe_edit", recipe_id=recipe.id))
    return render_template("recipes.html", recipes=Recipe.query.order_by(Recipe.id.desc()).all(), products=products)


@main.route("/recipes/<int:recipe_id>", methods=["GET", "POST"])
@manager_required
def recipe_edit(recipe_id):
    recipe = db.get_or_404(Recipe, recipe_id)
    if request.method == "POST":
        component = db.get_or_404(Product, request.form.get("component_id", type=int))
        if component.id == recipe.product_id:
            flash("Товар не может входить в собственную комплектацию", "danger")
        else:
            item = RecipeItem.query.filter_by(recipe_id=recipe.id, component_id=component.id).first() or RecipeItem(recipe_id=recipe.id, component_id=component.id)
            item.quantity = max(decimal_value(request.form.get("quantity")), Decimal("0.0001"))
            item.quantity_unit = normalize_unit(request.form.get("quantity_unit"))
            item.loss_percent = max(decimal_value(request.form.get("loss_percent")), 0)
            try:
                convert_quantity(item.quantity, item.quantity_unit, component.unit)
            except ValueError:
                flash("Для весового ингредиента выберите граммы или килограммы; для штучного — штуки", "danger")
            else:
                db.session.add(item)
                db.session.flush()
                recalculate_recipe(recipe)
                db.session.commit()
                flash("Компонент добавлен, себестоимость пересчитана", "success")
        return redirect(url_for("main.recipe_edit", recipe_id=recipe.id))
    recalculate_recipe(recipe)
    db.session.commit()
    return render_template("recipe_edit.html", recipe=recipe, components=Product.query.filter(Product.id != recipe.product_id, Product.is_active.is_(True)).order_by(Product.name).all())


@main.route("/production", methods=["GET", "POST"])
@manager_required
def production():
    finished = Product.query.filter_by(is_active=True, product_type="finished").order_by(Product.name).all()
    warehouses_rows = Warehouse.query.filter_by(is_active=True).order_by(Warehouse.name).all()
    if request.method == "POST":
        product = db.get_or_404(Product, request.form.get("product_id", type=int))
        warehouse = db.get_or_404(Warehouse, request.form.get("warehouse_id", type=int))
        source_warehouse = db.get_or_404(Warehouse, request.form.get("source_warehouse_id", type=int))
        quantity = decimal_value(request.form.get("planned_quantity"))
        if quantity <= 0 or not Recipe.query.filter_by(product_id=product.id).first():
            flash("Укажите количество и создайте комплектацию товара", "danger")
        else:
            number = f"PRD-{datetime.now():%Y%m%d}-{uuid4().hex[:5].upper()}"
            db.session.add(ProductionOrder(number=number, product_id=product.id, source_warehouse_id=source_warehouse.id, warehouse_id=warehouse.id, planned_quantity=quantity, planned_date=datetime.strptime(request.form.get("planned_date"), "%Y-%m-%d").date()))
            db.session.commit()
            flash("Производственное задание создано", "success")
        return redirect(url_for("main.production"))
    return render_template("production.html", orders=ProductionOrder.query.order_by(ProductionOrder.planned_date.desc(), ProductionOrder.id.desc()).all(), products=finished, warehouses=warehouses_rows)


@main.route("/production/<int:order_id>/complete", methods=["POST"])
@manager_required
def production_complete(order_id):
    production_order = db.get_or_404(ProductionOrder, order_id)
    if production_order.status == "completed":
        return redirect(url_for("main.production"))
    actual = decimal_value(request.form.get("actual_quantity"))
    recipe = Recipe.query.filter_by(product_id=production_order.product_id).first()
    if actual <= 0 or not recipe:
        flash("Укажите фактический выпуск", "danger")
        return redirect(url_for("main.production"))
    factor = actual / recipe.output_quantity
    requirements = []
    for item in recipe.items:
        required_per_recipe = convert_quantity(item.quantity, item.quantity_unit or item.component.unit, item.component.unit)
        required = required_per_recipe * factor * (Decimal("1") + item.loss_percent / Decimal("100"))
        balance = StockBalance.query.filter_by(warehouse_id=production_order.source_warehouse_id, product_id=item.component_id).first()
        if not balance or balance.quantity < required:
            available = balance.quantity if balance else 0
            flash(f"Недостаточно «{item.component.name}»: нужно {required}, доступно {available}", "danger")
            return redirect(url_for("main.production"))
        requirements.append((balance, required))
    for balance, required in requirements:
        balance.quantity -= required
        db.session.add(balance)
    finished_balance = StockBalance.query.filter_by(warehouse_id=production_order.warehouse_id, product_id=production_order.product_id).first() or StockBalance(warehouse_id=production_order.warehouse_id, product_id=production_order.product_id, quantity=0)
    finished_balance.quantity += actual
    production_order.actual_quantity = actual
    production_order.actual_cost = recalculate_recipe(recipe)
    production_order.status = "completed"
    production_order.produced_at = datetime.utcnow()
    production_order.batch_number = (request.form.get("batch_number") or production_order.number).strip()
    if production_order.product.shelf_life_days:
        production_order.expires_at = date.today() + timedelta(days=production_order.product.shelf_life_days)
    db.session.add(finished_balance)
    db.session.commit()
    flash("Выпуск проведён: сырьё списано, готовая продукция оприходована", "success")
    return redirect(url_for("main.production"))


@main.route("/mobile/returns/new/<int:order_id>", methods=["GET", "POST"])
@login_required
def mobile_return_new(order_id):
    order = db.get_or_404(Order, order_id)
    if not current_user.is_manager and order.representative_id != current_user.id:
        abort(403)
    if not order.customer.counterparty:
        flash("Для возврата привяжите торговую точку к контрагенту", "danger")
        return redirect(url_for("main.mobile"))
    warehouses_rows = Warehouse.query.filter_by(is_active=True).order_by(Warehouse.name).all()
    returned_rows = (
        db.session.query(SalesReturnItem.product_id, func.coalesce(func.sum(SalesReturnItem.quantity), 0))
        .join(SalesReturn)
        .filter(SalesReturn.order_id == order.id)
        .group_by(SalesReturnItem.product_id)
        .all()
    )
    returned = {product_id: quantity for product_id, quantity in returned_rows}
    available = {item.product_id: max(item.quantity - returned.get(item.product_id, 0), 0) for item in order.items}
    if request.method == "POST":
        warehouse = db.get_or_404(Warehouse, request.form.get("warehouse_id", type=int))
        selected, total = [], Decimal("0")
        for item in order.items:
            quantity = decimal_value(request.form.get(f"qty_{item.id}"))
            maximum = available[item.product_id]
            if quantity < 0 or quantity > maximum:
                flash(f"По товару «{item.product.name}» можно вернуть не более {maximum}", "danger")
                return redirect(url_for("main.mobile_return_new", order_id=order.id))
            if quantity > 0:
                selected.append((item, quantity))
                total += quantity * item.price
        if not selected:
            flash("Укажите количество возвращаемого товара", "danger")
        else:
            sales_return = SalesReturn(
                number=f"RET-{datetime.now():%Y%m%d}-{uuid4().hex[:5].upper()}",
                order_id=order.id,
                counterparty_id=order.customer.counterparty_id,
                warehouse_id=warehouse.id,
                reason=(request.form.get("reason") or "Возврат по заказу").strip(),
                status="draft",
                total_amount=total,
            )
            db.session.add(sales_return)
            db.session.flush()
            for item, quantity in selected:
                db.session.add(SalesReturnItem(return_id=sales_return.id, product_id=item.product_id, quantity=quantity, price=item.price))
            db.session.commit()
            flash(f"Возврат {sales_return.number} оформлен и передан руководителю", "success")
            return redirect(url_for("main.mobile"))
    return render_template("mobile_return.html", order=order, warehouses=warehouses_rows, available=available)


@main.route("/returns", methods=["GET", "POST"])
@manager_required
def returns():
    counterparties_rows = Counterparty.query.filter_by(is_active=True).order_by(Counterparty.name).all()
    warehouses_rows = Warehouse.query.filter_by(is_active=True).order_by(Warehouse.name).all()
    if request.method == "POST":
        counterparty = db.get_or_404(Counterparty, request.form.get("counterparty_id", type=int))
        warehouse = db.get_or_404(Warehouse, request.form.get("warehouse_id", type=int))
        number = f"RET-{datetime.now():%Y%m%d}-{uuid4().hex[:5].upper()}"
        row = SalesReturn(number=number, counterparty_id=counterparty.id, warehouse_id=warehouse.id, reason=(request.form.get("reason") or "").strip(), status="draft")
        db.session.add(row)
        db.session.commit()
        flash("Возврат создан", "success")
        return redirect(url_for("main.return_edit", return_id=row.id))
    return render_template("returns.html", returns=SalesReturn.query.order_by(SalesReturn.created_at.desc()).all(), counterparties=counterparties_rows, warehouses=warehouses_rows)


@main.route("/returns/<int:return_id>", methods=["GET", "POST"])
@manager_required
def return_edit(return_id):
    sales_return = db.get_or_404(SalesReturn, return_id)
    if request.method == "POST" and sales_return.status == "draft":
        product = db.get_or_404(Product, request.form.get("product_id", type=int))
        quantity = decimal_value(request.form.get("quantity"))
        price = decimal_value(request.form.get("price"), product.price)
        if quantity > 0:
            db.session.add(SalesReturnItem(return_id=sales_return.id, product_id=product.id, quantity=quantity, price=price))
            sales_return.total_amount += quantity * price
            db.session.commit()
        return redirect(url_for("main.return_edit", return_id=sales_return.id))
    return render_template("return_edit.html", sales_return=sales_return, products=Product.query.filter_by(is_active=True).order_by(Product.name).all())


@main.route("/returns/<int:return_id>/post", methods=["POST"])
@manager_required
def return_post(return_id):
    sales_return = db.get_or_404(SalesReturn, return_id)
    if sales_return.status == "draft" and sales_return.items:
        for item in sales_return.items:
            balance = StockBalance.query.filter_by(warehouse_id=sales_return.warehouse_id, product_id=item.product_id).first() or StockBalance(warehouse_id=sales_return.warehouse_id, product_id=item.product_id, quantity=0)
            balance.quantity += item.quantity
            db.session.add(balance)
        sales_return.status = "completed"
        db.session.commit()
        flash("Возврат проведён и учтён в остатках", "success")
    return redirect(url_for("main.returns"))


def order_payload(order):
    counterparty = order.customer.counterparty
    return {"number": order.number, "external_id": order.external_id, "status": order.status, "created_at": order.created_at.isoformat(), "customer": {"id": order.customer.id, "name": order.customer.name, "address": order.customer.address}, "counterparty": {"external_id": counterparty.external_id, "name": counterparty.name, "tax_id": counterparty.tax_id} if counterparty else None, "representative": {"id": order.representative.id, "name": order.representative.full_name}, "total_amount": float(order.total_amount), "items": [{"external_id": item.product.external_id, "barcode": item.product.barcode, "name": item.product.name, "quantity": float(item.quantity), "price": float(item.price)} for item in order.items]}


@api.route("/health")
@one_c_required
def api_health():
    return jsonify(ok=True, service="sinana")


@api.route("/products")
@one_c_required
def api_products():
    selected_group = setting("one_c_product_group_id")
    query = Product.query.filter_by(is_active=True)
    if selected_group:
        query = query.filter_by(external_group_id=selected_group)
    return jsonify(items=[{"external_id": p.external_id, "group_id": p.external_group_id, "barcode": p.barcode, "name": p.name, "unit": p.unit, "type": p.product_type, "price": float(p.price), "cost": float(p.calculated_cost), "stock": float(p.stock_qty)} for p in query.all()])


@api.route("/catalog", methods=["POST"])
@one_c_required
def api_catalog():
    payload = request.get_json(silent=True) or {}
    group_id = str(payload.get("group_id") or "").strip()
    selected_group = setting("one_c_product_group_id")
    if selected_group and group_id != selected_group:
        return jsonify(error="Передана не настроенная группа номенклатуры", expected_group_id=selected_group), 400
    rows = payload.get("items")
    if not isinstance(rows, list):
        return jsonify(error="Ожидается массив items"), 400
    updated = 0
    received_ids = set()
    for row in rows:
        external_id = str(row.get("external_id") or "").strip()
        barcode = barcode_value(row.get("barcode"))
        name = str(row.get("name") or "").strip()
        if not external_id or not barcode or not name:
            continue
        product = Product.query.filter((Product.external_id == external_id) | (Product.barcode == barcode)).first() or Product(barcode=barcode)
        product.external_id, product.external_group_id, product.barcode, product.name = external_id, group_id, barcode, name
        product.unit = normalize_unit(row.get("unit"))
        product.product_type = str(row.get("type") or "finished").strip() if str(row.get("type") or "finished") in {"raw", "semi_finished", "finished", "packaging"} else "finished"
        product.shelf_life_days = int(row["shelf_life_days"]) if row.get("shelf_life_days") not in (None, "") else None
        product.allergens = str(row.get("allergens") or "").strip() or None
        product.calculated_cost = decimal_value(row.get("cost"), product.calculated_cost)
        product.price = decimal_value(row.get("price"), product.price)
        product.is_active = True
        db.session.add(product)
        received_ids.add(external_id)
        updated += 1
    if payload.get("full_sync") and group_id:
        Product.query.filter(Product.external_group_id == group_id, Product.external_id.notin_(received_ids or {""})).update({Product.is_active: False}, synchronize_session=False)
    db.session.add(SyncLog(source="1c_catalog", status="success", items_count=updated, message=f"Группа {group_id}"))
    db.session.commit()
    return jsonify(ok=True, updated=updated, group_id=group_id)


@api.route("/counterparties", methods=["POST"])
@one_c_required
def api_counterparties():
    rows = (request.get_json(silent=True) or {}).get("items")
    if not isinstance(rows, list):
        return jsonify(error="Ожидается массив items"), 400
    updated = 0
    for item in rows:
        external_id = str(item.get("external_id") or "").strip()
        name = str(item.get("name") or "").strip()
        if not external_id or not name:
            continue
        row = Counterparty.query.filter_by(external_id=external_id).first() or Counterparty(external_id=external_id)
        row.name, row.tax_id = name, str(item.get("tax_id") or "").strip() or None
        row.registration_code = str(item.get("registration_code") or "").strip() or None
        row.legal_address, row.phone, row.email = str(item.get("address") or "").strip() or None, str(item.get("phone") or "").strip() or None, str(item.get("email") or "").strip() or None
        row.debt, row.is_active = decimal_value(item.get("debt")), bool(item.get("is_active", True))
        row.payment_terms_days = max(int(item.get("payment_terms_days") or 0), 0)
        price_type_external_id = str(item.get("price_type_external_id") or "").strip()
        if price_type_external_id:
            price_type = PriceType.query.filter_by(external_id=price_type_external_id).first()
            if price_type:
                row.price_type_id = price_type.id
        db.session.add(row)
        updated += 1
    db.session.add(SyncLog(source="1c_counterparties", status="success", items_count=updated))
    db.session.commit()
    return jsonify(ok=True, updated=updated)


@api.route("/prices", methods=["POST"])
@one_c_required
def api_prices():
    payload = request.get_json(silent=True) or {}
    types = payload.get("price_types", [])
    prices = payload.get("prices", [])
    if not isinstance(types, list) or not isinstance(prices, list):
        return jsonify(error="Ожидаются массивы price_types и prices"), 400
    for item in types:
        external_id, name = str(item.get("external_id") or "").strip(), str(item.get("name") or "").strip()
        if not external_id or not name:
            continue
        row = PriceType.query.filter_by(external_id=external_id).first() or PriceType(external_id=external_id, name=name)
        row.name, row.is_active = name, bool(item.get("is_active", True))
        db.session.add(row)
    db.session.flush()
    updated = 0
    for item in prices:
        product = Product.query.filter_by(external_id=str(item.get("product_external_id") or "").strip()).first()
        price_type = PriceType.query.filter_by(external_id=str(item.get("price_type_external_id") or "").strip()).first()
        if not product or not price_type:
            continue
        row = ProductPrice.query.filter_by(product_id=product.id, price_type_id=price_type.id).first() or ProductPrice(product_id=product.id, price_type_id=price_type.id)
        row.price = max(decimal_value(item.get("price")), 0)
        db.session.add(row)
        updated += 1
    db.session.add(SyncLog(source="1c_prices", status="success", items_count=updated))
    db.session.commit()
    return jsonify(ok=True, updated=updated, price_types=len(types))


@api.route("/warehouses", methods=["POST"])
@one_c_required
def api_warehouses():
    rows = (request.get_json(silent=True) or {}).get("items")
    if not isinstance(rows, list):
        return jsonify(error="Ожидается массив items"), 400
    updated = 0
    for item in rows:
        external_id, name = str(item.get("external_id") or "").strip(), str(item.get("name") or "").strip()
        if not external_id or not name:
            continue
        row = Warehouse.query.filter_by(external_id=external_id).first() or Warehouse(external_id=external_id, name=name)
        row.name, row.warehouse_type, row.is_active = name, str(item.get("type") or "finished"), bool(item.get("is_active", True))
        db.session.add(row)
        updated += 1
    db.session.commit()
    return jsonify(ok=True, updated=updated)


@api.route("/stocks", methods=["POST"])
@one_c_required
def api_stocks():
    items = (request.get_json(silent=True) or {}).get("items")
    if not isinstance(items, list):
        return jsonify(error="Ожидается массив items"), 400
    updated, unknown = 0, []
    for item in items:
        barcode = str(item.get("barcode") or "").strip()
        product = Product.query.filter_by(barcode=barcode).first()
        if not product:
            unknown.append(barcode)
            continue
        quantity = decimal_value(item.get("quantity"))
        warehouse_external_id = str(item.get("warehouse_external_id") or "").strip()
        warehouse = Warehouse.query.filter_by(external_id=warehouse_external_id).first() if warehouse_external_id else None
        if warehouse:
            balance = StockBalance.query.filter_by(warehouse_id=warehouse.id, product_id=product.id).first() or StockBalance(warehouse_id=warehouse.id, product_id=product.id)
            balance.quantity = quantity
            db.session.add(balance)
        else:
            product.stock_qty = quantity
        if item.get("price") is not None:
            product.price = decimal_value(item.get("price"))
        updated += 1
    db.session.flush()
    for product in Product.query.all():
        total = db.session.query(func.coalesce(func.sum(StockBalance.quantity), 0)).filter_by(product_id=product.id).scalar()
        if total:
            product.stock_qty = total
    db.session.add(SyncLog(source="1c_stock", status="success", items_count=updated, message=f"Не найдено: {len(unknown)}"))
    db.session.commit()
    return jsonify(ok=True, updated=updated, unknown_barcodes=unknown)


@api.route("/orders")
@one_c_required
def api_orders():
    query = Order.query.filter(Order.status != "cancelled")
    if request.args.get("unsynced", "1") != "0":
        query = query.filter(Order.synced_at.is_(None))
    return jsonify(items=[order_payload(order) for order in query.order_by(Order.created_at).limit(500).all()])


@api.route("/orders/ack", methods=["POST"])
@one_c_required
def api_orders_ack():
    rows = (request.get_json(silent=True) or {}).get("orders")
    if not isinstance(rows, list):
        return jsonify(error="Ожидается массив orders"), 400
    updated, missing = 0, []
    for row in rows:
        number = str(row.get("number") or "").strip()
        order = Order.query.filter_by(number=number).first()
        if not order:
            missing.append(number)
            continue
        external_id = str(row.get("external_id") or "").strip() or None
        if external_id and Order.query.filter(Order.external_id == external_id, Order.id != order.id).first():
            continue
        order.external_id = external_id or order.external_id
        order.external_number = str(row.get("external_number") or "").strip() or order.external_number
        status = str(row.get("status") or "").strip()
        if status in ORDER_STATUSES:
            order.status = status
        order.synced_at, order.sync_error = datetime.utcnow(), None
        updated += 1
    db.session.add(SyncLog(source="1c_orders", status="success", items_count=updated, message=f"Не найдено: {len(missing)}"))
    db.session.commit()
    return jsonify(ok=True, updated=updated, missing=missing)


@api.route("/orders/statuses", methods=["POST"])
@one_c_required
def api_order_statuses():
    return api_orders_ack()


@api.route("/returns")
@one_c_required
def api_returns():
    rows = SalesReturn.query.filter_by(status="completed", synced_at=None).order_by(SalesReturn.created_at).limit(500).all()
    return jsonify(items=[{"number": row.number, "external_id": row.external_id, "counterparty_external_id": row.counterparty.external_id, "warehouse_external_id": row.warehouse.external_id, "reason": row.reason, "total_amount": float(row.total_amount), "items": [{"barcode": item.product.barcode, "quantity": float(item.quantity), "price": float(item.price)} for item in row.items]} for row in rows])


@api.route("/returns/ack", methods=["POST"])
@one_c_required
def api_returns_ack():
    rows = (request.get_json(silent=True) or {}).get("returns")
    if not isinstance(rows, list):
        return jsonify(error="Ожидается массив returns"), 400
    updated = 0
    for item in rows:
        row = SalesReturn.query.filter_by(number=str(item.get("number") or "")).first()
        if row:
            row.external_id = str(item.get("external_id") or "").strip() or row.external_id
            row.synced_at = datetime.utcnow()
            updated += 1
    db.session.commit()
    return jsonify(ok=True, updated=updated)
